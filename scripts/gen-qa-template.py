"""Generates qa/templates/TestPlan.xlsx (run: python scripts/gen-qa-template.py).

The TestCases headers must match DEFAULT_COLUMNS in qa/src/config.ts; tests/qa-template.test.ts checks that.
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).resolve().parent.parent / "qa" / "templates" / "TestPlan.xlsx"

HEADER_FILL = PatternFill("solid", start_color="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF", name="Calibri")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
GREEN = PatternFill("solid", start_color="C6EFCE", end_color="C6EFCE")
RED = PatternFill("solid", start_color="FFC7CE", end_color="FFC7CE")
AMBER = PatternFill("solid", start_color="FFEB9C", end_color="FFEB9C")
GREY = PatternFill("solid", start_color="EDEDED", end_color="EDEDED")
WRAP = Alignment(wrap_text=True, vertical="top")

TEST_COLUMNS = [
    ("TC_ID", 12), ("Req_ID", 12), ("Module", 16), ("Title", 34), ("Preconditions", 30), ("Steps", 46), ("Expected", 40),
    ("Type", 13), ("Priority", 10), ("Role", 12), ("Automation_Status", 18), ("Script_Path", 40), ("Last_Run", 17),
    ("Last_Result", 13), ("Duration_s", 11), ("Run_Evidence", 36), ("Notes", 28),
]
COL = {name: get_column_letter(i + 1) for i, (name, _) in enumerate(TEST_COLUMNS)}
ROWS = 1000      # TestCases rows covered by formulas
REQ_ROWS = 200   # Requirements rows covered by the RTM


def header(sheet, names, widths=None, row=1):
    for i, name in enumerate(names, start=1):
        cell = sheet.cell(row=row, column=i, value=name)
        cell.fill, cell.font, cell.border = HEADER_FILL, HEADER_FONT, BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        if widths:
            sheet.column_dimensions[get_column_letter(i)].width = widths[i - 1]
    sheet.row_dimensions[row].height = 24
    sheet.freeze_panes = f"A{row + 1}"  # by address: touching cell A2 would push appended rows down


wb = Workbook()

# ------------------------------------------------------------------ README
readme = wb.active
readme.title = "README"
readme.column_dimensions["A"].width = 110
lines = [
    ("Test plan with automation", True),
    ("", False),
    ("1. Requirements: list each requirement with a Req_ID (REQ_001, REQ_002...).", False),
    ("2. TestCases: ask the assistant to write test cases from the requirements, or add rows by hand. Every row needs a TC_ID and a Req_ID.", False),
    ("3. Automation: the assistant writes a Playwright script per test case (saved as a draft in the QA project folder).", False),
    ("4. Approval: you read the script and run  excel-partner qa approve TC_001  . Only approved scripts run unattended.", False),
    ("5. Run: ask the assistant to run the tests, then write the results to this sheet (Last_Run, Last_Result, Duration_s, Run_Evidence).", False),
    ("6. RTM and Dashboard update by themselves from TestCases and Requirements.", False),
    ("", False),
    ("Columns filled by the tool (do not type in them): Automation_Status, Script_Path, Last_Run, Last_Result, Duration_s, Run_Evidence.", False),
    ("Never write passwords or tokens in this workbook. Sign-in is done once with  excel-partner qa login <role>  .", False),
    ("Rows 2-3 of TestCases are samples; delete or replace them.", False),
]
for i, (text, bold) in enumerate(lines, start=1):
    cell = readme.cell(row=i, column=1, value=text)
    cell.font = Font(bold=bold, size=14 if bold else 11)
    cell.alignment = Alignment(wrap_text=True, vertical="top")

# ------------------------------------------------------------------ Requirements
req = wb.create_sheet("Requirements")
req_cols = [("Req_ID", 12), ("Title", 36), ("Description", 60), ("Priority", 10), ("Source", 22), ("Status", 14)]
header(req, [c for c, _ in req_cols], [w for _, w in req_cols])
req.append(["REQ_001", "Users can sign in", "A registered user can sign in with a valid login and password and sees the home screen.", "High", "Sample", "Approved"])
req.append(["REQ_002", "Sale order can be confirmed", "A draft quotation with at least one line can be confirmed and becomes a sale order.", "High", "Sample", "Draft"])
for row in req.iter_rows(min_row=2, max_row=REQ_ROWS + 1, max_col=6):
    for cell in row:
        cell.alignment, cell.border = WRAP, BORDER
dv = DataValidation(type="list", formula1='"High,Medium,Low"', allow_blank=True)
req.add_data_validation(dv)
dv.add(f"D2:D{REQ_ROWS + 1}")

# ------------------------------------------------------------------ TestCases
tc = wb.create_sheet("TestCases")
header(tc, [c for c, _ in TEST_COLUMNS], [w for _, w in TEST_COLUMNS])
tc.append(["TC_001", "REQ_001", "Login", "Sign in with a valid account", "Account exists and is active", "1. Open the login page\n2. Enter a valid login and password\n3. Press Log in", "The home screen opens and the user's name is shown", "Positive", "High", "admin", None, None, None, None, None, None, "Sample"])
tc.append(["TC_002", "REQ_001", "Login", "Sign in with a wrong password", "Account exists", "1. Open the login page\n2. Enter a valid login and a wrong password\n3. Press Log in", "An error 'Wrong login/password' is shown and the user stays on the login page", "Negative", "High", "admin", None, None, None, None, None, None, "Sample"])
for row in tc.iter_rows(min_row=2, max_row=ROWS + 1, max_col=len(TEST_COLUMNS)):
    for cell in row:
        cell.alignment, cell.border = WRAP, BORDER
for r in range(2, ROWS + 2):
    tc[f"{COL['Duration_s']}{r}"].number_format = "0.0"
for name in ("Automation_Status", "Script_Path", "Last_Run", "Last_Result", "Duration_s", "Run_Evidence"):
    tc[f"{COL[name]}1"].fill = PatternFill("solid", start_color="7F6000")  # tool-filled columns
    for r in range(2, ROWS + 2):
        tc[f"{COL[name]}{r}"].fill = GREY

dv_type = DataValidation(type="list", formula1='"Positive,Negative,Boundary,Permission,Regression,Usability"', allow_blank=True)
dv_prio = DataValidation(type="list", formula1='"High,Medium,Low"', allow_blank=True)
dv_req = DataValidation(type="list", formula1=f"=Requirements!$A$2:$A${REQ_ROWS + 1}", allow_blank=True)
dv_req.errorTitle, dv_req.error = "Unknown requirement", "Pick a Req_ID from the Requirements sheet."
for validation in (dv_type, dv_prio, dv_req):
    tc.add_data_validation(validation)
dv_type.add(f"{COL['Type']}2:{COL['Type']}{ROWS + 1}")
dv_prio.add(f"{COL['Priority']}2:{COL['Priority']}{ROWS + 1}")
dv_req.add(f"{COL['Req_ID']}2:{COL['Req_ID']}{ROWS + 1}")

result_range = f"{COL['Last_Result']}2:{COL['Last_Result']}{ROWS + 1}"
tc.conditional_formatting.add(result_range, CellIsRule(operator="equal", formula=['"Passed"'], fill=GREEN))
tc.conditional_formatting.add(result_range, CellIsRule(operator="equal", formula=['"Failed"'], fill=RED))
tc.conditional_formatting.add(result_range, CellIsRule(operator="equal", formula=['"Timed out"'], fill=RED))
tc.conditional_formatting.add(result_range, CellIsRule(operator="equal", formula=['"Flaky"'], fill=AMBER))
status_range = f"{COL['Automation_Status']}2:{COL['Automation_Status']}{ROWS + 1}"
tc.conditional_formatting.add(status_range, CellIsRule(operator="equal", formula=['"Approved"'], fill=GREEN))
tc.conditional_formatting.add(status_range, CellIsRule(operator="equal", formula=['"Modified"'], fill=AMBER))
tc.conditional_formatting.add(status_range, CellIsRule(operator="equal", formula=['"Invalid"'], fill=RED))
tc.auto_filter.ref = f"A1:{get_column_letter(len(TEST_COLUMNS))}{ROWS + 1}"

# ------------------------------------------------------------------ RTM
rtm = wb.create_sheet("RTM")
rtm_cols = [("Req_ID", 12), ("Requirement", 36), ("Test cases", 12), ("Automated (approved)", 20), ("Passed", 10), ("Failed", 10), ("Not run", 10), ("Coverage", 16)]
header(rtm, [c for c, _ in rtm_cols], [w for _, w in rtm_cols])
rid, rst, rres = COL["Req_ID"], COL["Automation_Status"], COL["Last_Result"]
for r in range(2, REQ_ROWS + 2):
    rtm[f"A{r}"] = f'=IF(Requirements!A{r}="","",Requirements!A{r})'
    rtm[f"B{r}"] = f'=IF(A{r}="","",Requirements!B{r})'
    rtm[f"C{r}"] = f'=IF(A{r}="","",COUNTIF(TestCases!${rid}$2:${rid}${ROWS + 1},A{r}))'
    rtm[f"D{r}"] = f'=IF(A{r}="","",COUNTIFS(TestCases!${rid}$2:${rid}${ROWS + 1},A{r},TestCases!${rst}$2:${rst}${ROWS + 1},"Approved"))'
    rtm[f"E{r}"] = f'=IF(A{r}="","",COUNTIFS(TestCases!${rid}$2:${rid}${ROWS + 1},A{r},TestCases!${rres}$2:${rres}${ROWS + 1},"Passed"))'
    rtm[f"F{r}"] = f'=IF(A{r}="","",COUNTIFS(TestCases!${rid}$2:${rid}${ROWS + 1},A{r},TestCases!${rres}$2:${rres}${ROWS + 1},"Failed")+COUNTIFS(TestCases!${rid}$2:${rid}${ROWS + 1},A{r},TestCases!${rres}$2:${rres}${ROWS + 1},"Timed out"))'
    rtm[f"G{r}"] = f'=IF(A{r}="","",C{r}-E{r}-F{r}-COUNTIFS(TestCases!${rid}$2:${rid}${ROWS + 1},A{r},TestCases!${rres}$2:${rres}${ROWS + 1},"Flaky"))'
    rtm[f"H{r}"] = f'=IF(A{r}="","",IF(C{r}=0,"No tests",IF(F{r}>0,"Failing",IF(E{r}=C{r},"Verified","In progress"))))'
    for c in range(1, 9):
        rtm.cell(row=r, column=c).border = BORDER
rng = f"H2:H{REQ_ROWS + 1}"
rtm.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"No tests"'], fill=RED))
rtm.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Failing"'], fill=RED))
rtm.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Verified"'], fill=GREEN))
rtm.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"In progress"'], fill=AMBER))

# ------------------------------------------------------------------ Dashboard
dash = wb.create_sheet("Dashboard")
dash.column_dimensions["A"].width = 34
dash.column_dimensions["B"].width = 16
dash["A1"] = "Test status"
dash["A1"].font = Font(bold=True, size=16)
res = f"TestCases!${rres}$2:${rres}${ROWS + 1}"
ids = f"TestCases!${COL['TC_ID']}$2:${COL['TC_ID']}${ROWS + 1}"
stat = f"TestCases!${rst}$2:${rst}${ROWS + 1}"
kpis = [
    ("Test cases", f'=COUNTA({ids})', "0"),
    ("Automated and approved", f'=COUNTIF({stat},"Approved")', "0"),
    ("Passed", f'=COUNTIF({res},"Passed")', "0"),
    ("Failed", f'=COUNTIF({res},"Failed")+COUNTIF({res},"Timed out")', "0"),
    ("Flaky", f'=COUNTIF({res},"Flaky")', "0"),
    ("Not run yet", '=B3-B5-B6-B7-COUNTIF(' + res + ',"Skipped")', "0"),
    ("Pass rate (of executed)", '=IF(B5+B6+B7=0,"-",B5/(B5+B6+B7))', "0.0%"),
    ("Requirements with no test", f'=COUNTIF(RTM!$H$2:$H${REQ_ROWS + 1},"No tests")', "0"),
    ("Requirement coverage", f'=IF(COUNTIF(Requirements!$A$2:$A${REQ_ROWS + 1},"?*")=0,"-",1-B10/COUNTIF(Requirements!$A$2:$A${REQ_ROWS + 1},"?*"))', "0.0%"),
]
for i, (label, formula, fmt) in enumerate(kpis, start=3):
    dash[f"A{i}"] = label
    dash[f"B{i}"] = formula
    dash[f"B{i}"].number_format = fmt
    dash[f"A{i}"].border = dash[f"B{i}"].border = BORDER
    dash[f"B{i}"].font = Font(bold=True)
dash.conditional_formatting.add("B6", CellIsRule(operator="greaterThan", formula=["0"], fill=RED))
dash.conditional_formatting.add("B10", CellIsRule(operator="greaterThan", formula=["0"], fill=RED))

dash["D2"] = "Result"
dash["E2"] = "Count"
for cell in (dash["D2"], dash["E2"]):
    cell.fill, cell.font, cell.border = HEADER_FILL, HEADER_FONT, BORDER
for i, (label, ref) in enumerate([("Passed", "B5"), ("Failed", "B6"), ("Flaky", "B7"), ("Not run", "B8")], start=3):
    dash[f"D{i}"] = label
    dash[f"E{i}"] = f"={ref}"
chart = BarChart()
chart.type, chart.style, chart.title = "bar", 10, "Results"
chart.y_axis.title, chart.legend = None, None
chart.add_data(Reference(dash, min_col=5, min_row=2, max_row=6), titles_from_data=True)
chart.set_categories(Reference(dash, min_col=4, min_row=3, max_row=6))
chart.height, chart.width = 7, 12
dash.add_chart(chart, "G2")

# ------------------------------------------------------------------ Bugs
bugs = wb.create_sheet("Bugs")
bug_cols = [("Bug_ID", 11), ("TC_ID", 12), ("Title", 36), ("Severity", 11), ("Status", 12), ("Steps", 44), ("Expected", 30), ("Actual", 30), ("Evidence", 34), ("Reported", 12)]
header(bugs, [c for c, _ in bug_cols], [w for _, w in bug_cols])
dv_sev = DataValidation(type="list", formula1='"Critical,Major,Minor,Trivial"', allow_blank=True)
dv_bst = DataValidation(type="list", formula1='"New,Confirmed,Fixed,Verified,Rejected"', allow_blank=True)
bugs.add_data_validation(dv_sev)
bugs.add_data_validation(dv_bst)
dv_sev.add("D2:D500")
dv_bst.add("E2:E500")
for row in bugs.iter_rows(min_row=2, max_row=500, max_col=10):
    for cell in row:
        cell.alignment, cell.border = WRAP, BORDER

wb.move_sheet("Dashboard", offset=-3)  # README, Dashboard, Requirements, TestCases, RTM, Bugs
OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print(f"wrote {OUT} ({OUT.stat().st_size} bytes); sheets: {wb.sheetnames}")
